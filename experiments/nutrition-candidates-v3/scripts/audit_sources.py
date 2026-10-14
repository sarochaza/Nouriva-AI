"""Reconcile curated references against local original USDA ZIPs and optional XLSX."""
import argparse,collections,csv,io,json,math,sqlite3,zipfile
from pathlib import Path
from contextlib import closing
from src.common import ROOT,read_csv,sha256,save_json

FIELDS={'kcal_100g':'1008','protein_g_100g':'1003','fat_g_100g':'1004','carbs_g_100g':'1005'}

def zip_rows(z,suffix):
    names=[n for n in z.namelist() if n.rsplit('/',1)[-1]==suffix]
    if len(names)!=1:raise ValueError('Missing/ambiguous USDA table '+suffix)
    with z.open(names[0]) as stream:
        yield from csv.DictReader(io.TextIOWrapper(stream,encoding='utf-8-sig'))

def audit_sources(sources,xlsx=None):
    sources=Path(sources);rows=read_csv(sources/'nouriva_usda_food_nutrition.csv')
    mismatches=[];raw_audits=[];checks=0;source_checked=set();ndb={};rounding=[]
    for path in sorted(sources.glob('FoodData_Central_*.zip')):
        source='SR_Legacy' if 'sr_legacy' in path.name else ('Foundation' if 'foundation' in path.name else 'FNDDS')
        selected={r['fdc_id']:r for r in rows if r['source']==source}
        if not selected:continue
        with zipfile.ZipFile(path) as z:
            raw_food={r['fdc_id']:r for r in zip_rows(z,'food.csv')}
            nutrients={r['id']:r for r in zip_rows(z,'nutrient.csv')}
            raw_values=collections.defaultdict(list);raw_count=0;identifier_counts=collections.Counter()
            for r in zip_rows(z,'food_nutrient.csv'):
                raw_count+=1
                identifier_counts[r['nutrient_id']]+=1
                if r['fdc_id'] in selected and r['nutrient_id'] in ('1003','1004','1005','1008','2047','2048','203','204','205','208') and r['amount'].strip():
                    raw_values[(r['fdc_id'],r['nutrient_id'])].append(float(r['amount']))
            new_count=sum(identifier_counts[x] for x in ('1003','1004','1005','1008'))
            old_count=sum(identifier_counts[x] for x in ('203','204','205','208'))
            domain='legacy_nutrient_number' if old_count and not new_count else 'fdc_nutrient_id'
            if old_count and new_count:raise ValueError('Mixed nutrient identifier domains in '+path.name)
            old_ids={'1003':'203','1004':'204','1005':'205','1008':'208'}
            for fid,r in selected.items():
                food=raw_food.get(fid)
                if not food or food['description']!=r['name_en']:mismatches.append({'source':source,'fdc_id':fid,'field':'name_en','reason':'missing/mismatched original description'})
                for field,nid in FIELDS.items():
                    if field=='kcal_100g' and r.get('energy_method','').endswith('2048'):nid='2048'
                    elif field=='kcal_100g' and r.get('energy_method','').endswith('2047'):nid='2047'
                    expected_unit='KCAL' if field=='kcal_100g' else 'G'
                    raw_unit=nutrients.get(nid,{}).get('unit_name','').upper()
                    if domain=='legacy_nutrient_number':
                        nid=old_ids[nid]
                        unit_rows=[n for n in nutrients.values() if n.get('nutrient_nbr')==nid]
                        if len(unit_rows)!=1:raise ValueError('Legacy nutrient number has no unique dictionary entry')
                        raw_unit=unit_rows[0]['unit_name'].upper()
                    if raw_unit!=expected_unit:mismatches.append({'source':source,'fdc_id':fid,'field':field,'reason':'unexpected original nutrient unit','unit':raw_unit})
                    if r.get(field) in ('',None):continue
                    checks+=1;curated=float(r[field]);values=raw_values.get((fid,nid),[])
                    if not any(math.isclose(curated,v,rel_tol=1e-7,abs_tol=1e-6) for v in values):
                        mismatches.append({'source':source,'fdc_id':fid,'field':field,'curated_value':curated,'raw_values':values[:10]})
                    elif values and not any(curated==v for v in values):
                        rounding.append({'source':source,'fdc_id':fid,'field':field,'curated_value':curated,'raw_values':values[:3]})
            if source=='SR_Legacy' and any(n.endswith('/sr_legacy_food.csv') for n in z.namelist()):
                ndb={r['NDB_number'].zfill(5):'SR_Legacy:2018-04:'+r['fdc_id'] for r in zip_rows(z,'sr_legacy_food.csv')}
            raw_audits.append({'source':source,'filename':path.name,'sha256':sha256(path),'raw_food_rows':len(raw_food),'raw_food_nutrient_rows':raw_count,'curated_food_rows':len(selected),'nutrient_id_domain':domain,
                               'raw_data_type_counts':dict(collections.Counter(r['data_type'] for r in raw_food.values()))})
            source_checked.add(source)
    original=sources/'original_nutrition.sqlite';sqlite_mismatches=[];sqlite_info={}
    if original.exists():
        with closing(sqlite3.connect(original)) as c:
            c.row_factory=sqlite3.Row
            database={r['fdc_id']:dict(r) for r in c.execute('SELECT * FROM food')}
            sqlite_info={'food_rows':len(database),'schema':[dict(r) for r in c.execute('PRAGMA table_info(food)')],'integrity':c.execute('PRAGMA integrity_check').fetchone()[0]}
            for r in rows:
                old=database.get(r['fdc_id'])
                if old is None:sqlite_mismatches.append({'fdc_id':r['fdc_id'],'reason':'ID not in original SQLite'});continue
                for k in FIELDS:
                    cur=None if r[k]=='' else float(r[k]);db=old[k]
                    if cur!=db:sqlite_mismatches.append({'fdc_id':r['fdc_id'],'field':k,'csv':cur,'sqlite':db})
    xlsx_report=None
    if xlsx:
        # Only the optional audit command needs openpyxl. Inference/evaluation use stdlib.
        from openpyxl import load_workbook
        w=load_workbook(xlsx,read_only=True,data_only=True);s=w['AUSNUT 2023'];iterator=s.iter_rows(values_only=True)
        next(iterator);next(iterator);header=[str(v).strip() if v is not None else '' for v in next(iterator)[:11]]
        original_measures={str(int(row[3])):dict(zip(header,row[:11])) for row in iterator if row[3] is not None}
        measure_mismatches=[];csv_measures=read_csv(sources/'ausnut_2023_food_measures_only.csv')
        for r in csv_measures:
            x=original_measures.get(r['Measure ID'])
            if not x:measure_mismatches.append({'id':r['Measure ID'],'reason':'missing measure'});continue
            for k in ('Quantity','Gram amount','Volume'):
                a=None if r[k]=='' else float(r[k]);b=x[k]
                if not (a is None and b is None) and (a is None or b is None or not math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-10)):measure_mismatches.append({'id':r['Measure ID'],'field':k})
            if r['Public food key']!=x['Public food key'] or r['Food name']!=x['Food name']:measure_mismatches.append({'id':r['Measure ID'],'field':'identity'})
        w.close();xlsx_report={'xlsx_sha256':sha256(xlsx),'original_rows':len(original_measures),'csv_rows':len(csv_measures),'mismatches':measure_mismatches}
    manifest=sources/'source_manifest.json';hash_audit=[]
    if manifest.exists():
        for r in json.loads(manifest.read_text()):
            name=Path(r['file']).name;path=sources/name
            if not path.exists() and xlsx and name.endswith('.xlsx'):path=Path(xlsx)
            hash_audit.append({'filename':name,'expected_sha256':r['sha256'],'observed_sha256':sha256(path) if path.exists() else None,'verified':path.exists() and sha256(path)==r['sha256']})
    return {'curated_rows':len(rows),'original_archives':raw_audits,'sources_checked':sorted(source_checked),'raw_nutrient_value_checks':checks,'mismatches':mismatches,
            'original_sqlite':sqlite_info,'original_sqlite_mismatches':sqlite_mismatches,'ausnut_xlsx_reconciliation':xlsx_report,'manifest_hash_checks':hash_audit,'ndb_crosswalk':ndb,
            'unverified_sources':sorted({r['source'] for r in rows}-source_checked),
            'rounding_differences_within_tolerance':rounding,'numeric_tolerance':{'relative':1e-7,'absolute':1e-6}}

def main():
    p=argparse.ArgumentParser();p.add_argument('--sources',type=Path,default=ROOT/'nutrition_database/sources');p.add_argument('--xlsx',type=Path);a=p.parse_args()
    audit=audit_sources(a.sources,a.xlsx);crosswalk=audit.pop('ndb_crosswalk');save_json(a.sources/'ndb_crosswalk.json',crosswalk)
    save_json(ROOT/'nutrition_database/normalized/source_integrity_audit.json',audit)
    print(json.dumps({'curated_rows':audit['curated_rows'],'raw_nutrient_value_checks':audit['raw_nutrient_value_checks'],'mismatch_count':len(audit['mismatches']),
                      'sqlite_mismatch_count':len(audit['original_sqlite_mismatches']),'rounding_differences':len(audit['rounding_differences_within_tolerance']),'unverified_sources':audit['unverified_sources']},indent=2))
    if audit['mismatches'] or audit['original_sqlite_mismatches'] or audit['unverified_sources'] or (audit['ausnut_xlsx_reconciliation'] and audit['ausnut_xlsx_reconciliation']['mismatches']):raise SystemExit(1)
if __name__=='__main__':main()
